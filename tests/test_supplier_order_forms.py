from __future__ import annotations

from io import BytesIO
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd

APP_DIR = Path(__file__).resolve().parents[1] / "app"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

from supplier_order_forms import (  # noqa: E402
    WQ_ORDER_ROWS,
    build_order_form_rows,
    apply_order_file,
    build_wq_order_workbook,
    order_form_from_items,
    order_form_supplier,
    order_items_from_form,
    parse_order_form_file,
)


def forecast(*rows: tuple[str, float]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "product": product,
                "recommended_order_qty": quantity,
                "abc_class": "A",
                "xyz_class": "X",
                "priority": "Высокий",
                "stock_status": "Нужно пополнение",
                "available_stock_qty": 3.0,
                "forecast_qty": 40.0,
            }
            for product, quantity in rows
        ]
    )


class OrderFormRowsTest(unittest.TestCase):
    def row(self, frame: pd.DataFrame, name: str) -> pd.Series:
        return frame[frame["name"] == name].iloc[0]

    def test_board_and_edge_of_one_decor_go_to_their_columns(self) -> None:
        result = build_order_form_rows(
            forecast(
                ("US77301/PGR4201 Маффин 18*1220*2800мм", 219.4),
                ("PGR4201 Маффин 1*22мм", 1500),
            )
        )
        muffin = self.row(result, "Маффин")
        self.assertEqual(muffin["quantity"], 220)
        self.assertEqual(muffin["edge_or_sheets"], 1500)
        self.assertEqual(muffin["abc_class"], "A")
        self.assertEqual(muffin["code"], "WQ63893M")

    def test_not_needed_position_stays_empty(self) -> None:
        result = build_order_form_rows(forecast(("US77306/PGR4206 Бисквит 18*1220*2800мм", 0)))
        biscuit = self.row(result, "Бисквит")
        self.assertEqual(biscuit["quantity"], 0)
        self.assertTrue(biscuit["found"])
        self.assertFalse(self.row(result, "Орео")["found"])

    def test_oreo_and_tiramisu_use_the_site_article(self) -> None:
        result = build_order_form_rows(
            forecast(
                ("US77304/PGR4204 Орео 18*1220*2800мм", 20),
                ("US77305/PGR4205 Тирамису 18*1220*2800мм", 40),
            )
        )
        self.assertEqual(self.row(result, "Орео")["code"], "US77304/PGR4204")
        self.assertEqual(self.row(result, "Орео")["quantity"], 20)
        self.assertEqual(self.row(result, "Тирамису")["code"], "US77305/PGR4205")
        self.assertEqual(self.row(result, "Тирамису")["quantity"], 40)

    def test_new_position_is_found_by_name_only_with_the_board_size(self) -> None:
        result = build_order_form_rows(
            forecast(
                ("Панель Титан 18*1220*2800мм", 20),
                ("Титан крепёж для панели", 500),
            )
        )
        self.assertEqual(self.row(result, "Титан")["quantity"], 20)

    def test_hpl_sheets_go_to_the_sheets_column_and_stay_out_of_the_boards_total(self) -> None:
        result = build_order_form_rows(
            forecast(
                ("HPL Compact WQS31008-MT Изысканный Камень", 5),
                ("HPL Compact белый 3050", 9),
            )
        )
        stone = self.row(result, "Изысканный Камень")
        self.assertEqual(stone["edge_or_sheets"], 5)
        self.assertEqual(stone["quantity"], 0)
        self.assertEqual(self.row(result, "белый")["edge_or_sheets"], 0)
        self.assertEqual(result["quantity"].sum(), 0)

    def test_empty_forecast_keeps_the_whole_form(self) -> None:
        result = build_order_form_rows(pd.DataFrame())
        self.assertEqual(len(result), len(WQ_ORDER_ROWS))
        self.assertEqual(result["quantity"].sum(), 0)


class OrderFromFormTest(unittest.TestCase):
    def setUp(self) -> None:
        self.rows = build_order_form_rows(
            forecast(
                ("US77301/PGR4201 Маффин 18*1220*2800мм", 220),
                ("PGR4201 Маффин 1*22мм", 1500),
                ("HPL Compact WQS1837-3 FR Молочный Кварцит", 10),
                ("US77306/PGR4206 Бисквит 18*1220*2800мм", 0),
            )
        )

    def test_order_lines_keep_the_site_products_and_skip_empty_rows(self) -> None:
        items = order_items_from_form(self.rows).set_index("product")["quantity"]
        self.assertEqual(
            items.to_dict(),
            {
                "US77301/PGR4201 Маффин 18*1220*2800мм": 220.0,
                "PGR4201 Маффин 1*22мм": 1500.0,
                "HPL Compact WQS1837-3 FR Молочный Кварцит": 10.0,
            },
        )

    def test_quantities_changed_in_the_form_reach_the_order(self) -> None:
        edited = self.rows.copy()
        edited.loc[edited["name"] == "Маффин", "quantity"] = 240
        items = order_items_from_form(edited).set_index("product")["quantity"]
        self.assertEqual(items["US77301/PGR4201 Маффин 18*1220*2800мм"], 240.0)

    def test_saved_order_rebuilds_the_same_form(self) -> None:
        rebuilt = order_form_from_items(order_items_from_form(self.rows))
        for column in ("code", "quantity", "edge_or_sheets"):
            self.assertEqual(rebuilt[column].tolist(), self.rows[column].tolist())

    def test_order_supplier_comes_from_the_site_or_falls_back(self) -> None:
        self.assertEqual(order_form_supplier(self.rows), "WQ")
        named = self.rows.assign(supplier="Кроношпан Китай")
        self.assertEqual(order_form_supplier(named), "Кроношпан Китай")


class UploadedOrderFileTest(unittest.TestCase):
    def setUp(self) -> None:
        self.rows = build_order_form_rows(
            forecast(
                ("US77301/PGR4201 Маффин 18*1220*2800мм", 220),
                ("PGR4201 Маффин 1*22мм", 1500),
                ("US77311/PGR4211 Чизкейк 18*1220*2800мм", 80),
                ("HPL Compact WQS1837-3 FR Молочный Кварцит", 10),
            )
        )
        self.workbook = load_workbook(BytesIO(build_wq_order_workbook(self.rows)))
        self.sheet = self.workbook.active

    def cell_row(self, name: str) -> int:
        for row in range(5, self.sheet.max_row + 1):
            if self.sheet.cell(row, 6).value == name:
                return row
        self.fail(f"{name} is not in the form")

    def content(self) -> bytes:
        buffer = BytesIO()
        self.workbook.save(buffer)
        return buffer.getvalue()

    def test_unchanged_file_gives_back_the_same_quantities(self) -> None:
        parsed = parse_order_form_file(self.content())
        self.assertEqual(parsed.errors, [])
        applied = apply_order_file(self.rows, parsed)
        self.assertEqual(applied["quantity"].tolist(), self.rows["quantity"].tolist())
        self.assertEqual(applied["edge_or_sheets"].tolist(), self.rows["edge_or_sheets"].tolist())

    def test_changed_and_cleared_quantities_become_final(self) -> None:
        self.sheet.cell(self.cell_row("Маффин"), 3).value = 200
        self.sheet.cell(self.cell_row("Чизкейк"), 3).value = None
        applied = apply_order_file(self.rows, parse_order_form_file(self.content()))
        items = order_items_from_form(applied).set_index("product")["quantity"]
        self.assertEqual(items["US77301/PGR4201 Маффин 18*1220*2800мм"], 200.0)
        self.assertNotIn("US77311/PGR4211 Чизкейк 18*1220*2800мм", items.index)

    def test_position_without_sales_keeps_its_quantity_and_survives_a_saved_order(self) -> None:
        self.sheet.cell(self.cell_row("Титан"), 3).value = 20
        applied = apply_order_file(self.rows, parse_order_form_file(self.content()))
        items = order_items_from_form(applied)
        self.assertIn("WQ63142LJ Титан 18*1220*2800мм", items["product"].tolist())
        rebuilt = order_form_from_items(items)
        self.assertEqual(rebuilt[rebuilt["name"] == "Титан"]["quantity"].iloc[0], 20)

    def test_hpl_sheets_are_read_from_either_column(self) -> None:
        self.sheet.cell(self.cell_row("Молочный Кварцит"), 3).value = 7
        self.sheet.cell(self.cell_row("Молочный Кварцит"), 5).value = None
        parsed = parse_order_form_file(self.content())
        self.assertEqual(parsed.quantities["WQS1837-3 FR"], (0, 7))

    def test_code_is_matched_despite_spaces_and_number_type(self) -> None:
        self.sheet.cell(self.cell_row("Маффин"), 2).value = " WQ63893M "
        self.sheet.cell(self.cell_row("белый"), 2).value = 81050
        self.sheet.cell(self.cell_row("белый"), 5).value = 5
        parsed = parse_order_form_file(self.content())
        self.assertEqual(parsed.quantities["WQ63893M"][0], 220)
        self.assertEqual(parsed.quantities["81050"], (0, 5))

    def test_unknown_code_and_bad_number_are_reported(self) -> None:
        extra = self.sheet.max_row + 1
        self.sheet.cell(extra, 2).value = "WQ999"
        self.sheet.cell(extra, 3).value = 40
        self.sheet.cell(self.cell_row("Маффин"), 3).value = "нет"
        parsed = parse_order_form_file(self.content())
        self.assertEqual(parsed.unknown_codes, ["WQ999"])
        self.assertEqual(len(parsed.errors), 1)
        self.assertIn("WQ63893M", parsed.errors[0])

    def test_not_a_form_is_rejected(self) -> None:
        self.assertTrue(parse_order_form_file(b"not an excel file").errors)
        other = BytesIO()
        workbook = load_workbook(BytesIO(build_wq_order_workbook(self.rows)))
        workbook.active.delete_rows(4)
        workbook.save(other)
        self.assertTrue(parse_order_form_file(other.getvalue()).errors)

    def test_fingerprint_identifies_the_same_file(self) -> None:
        content = self.content()
        self.assertEqual(parse_order_form_file(content).fingerprint, parse_order_form_file(content).fingerprint)
        self.sheet.cell(self.cell_row("Маффин"), 3).value = 1
        self.assertNotEqual(parse_order_form_file(content).fingerprint, parse_order_form_file(self.content()).fingerprint)


class OrderWorkbookTest(unittest.TestCase):
    def setUp(self) -> None:
        rows = build_order_form_rows(
            forecast(
                ("US77301/PGR4201 Маффин 18*1220*2800мм", 220),
                ("US77311/PGR4211 Чизкейк 18*1220*2800мм", 80),
                ("PGR4201 Маффин 1*22мм", 1500),
                ("HPL Compact WQS1837-3 FR Молочный Кварцит", 10),
            )
        )
        content = build_wq_order_workbook(rows)
        self.sheet = load_workbook(BytesIO(content)).active
        self.values = load_workbook(BytesIO(content), data_only=True).active

    def find_row(self, name: str) -> int:
        for row in range(5, self.sheet.max_row + 1):
            if self.sheet.cell(row, 6).value == name:
                return row
        self.fail(f"{name} is not in the form")

    def test_layout_matches_the_supplier_form(self) -> None:
        self.assertEqual(self.sheet["C3"].value, "2800*1220*18мм")
        self.assertEqual(
            [self.sheet.cell(4, column).value for column in range(2, 7)],
            ["Код", "кол-во", "код старый", "кромка", "Название"],
        )

    def test_quantities_are_written_and_empty_rows_stay_blank(self) -> None:
        muffin = self.find_row("Маффин")
        self.assertEqual(self.sheet.cell(muffin, 3).value, 220)
        self.assertEqual(self.sheet.cell(muffin, 5).value, 1500)
        self.assertIsNone(self.sheet.cell(self.find_row("Брауни"), 3).value)
        self.assertEqual(self.sheet.cell(self.find_row("Молочный Кварцит"), 5).value, 10)
        self.assertIsNone(self.sheet.cell(self.find_row("Молочный Кварцит"), 3).value)

    def test_only_the_boards_total_is_kept(self) -> None:
        total_row = self.sheet.max_row
        self.assertEqual(self.sheet.cell(total_row, 3).value, f"=SUM(C5:C{total_row - 1})")
        self.assertEqual(self.values.cell(total_row, 3).value, 300)
        for column in (2, 4, 5, 6):
            self.assertIsNone(self.sheet.cell(total_row, column).value)


if __name__ == "__main__":
    unittest.main()
