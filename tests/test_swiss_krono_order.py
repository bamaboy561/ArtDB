from __future__ import annotations

import sys
import tempfile
import unittest
import warnings
from io import BytesIO
from pathlib import Path
from unittest import mock
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook

APP_DIR = Path(__file__).resolve().parents[1] / "app"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

import swiss_krono_order as krono  # noqa: E402

# The factory's form is not in the repository; these tests run where a copy is at hand.
PRIVATE_TEMPLATE = Path(__file__).resolve().parent / "private" / "swiss_krono_ldsp.xlsx"
krono.TEMPLATE_PATH = PRIVATE_TEMPLATE


def setUpModule() -> None:
    if not PRIVATE_TEMPLATE.is_file():
        raise unittest.SkipTest("Нет копии бланка завода в tests/private")


def forecast(*rows: tuple) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"product": product, "supplier": supplier, "recommended_order_qty": qty, "priority": priority}
            for product, supplier, qty, priority in rows
        ]
    )


def workbook(content: bytes, *, values: bool = True):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(BytesIO(content), data_only=values)


class ParseProductTest(unittest.TestCase):
    def check(self, name: str, decor: str, structure: str, thickness: int, length: int) -> None:
        parsed = krono.parse_product(name)
        self.assertIsNotNone(parsed, name)
        self.assertEqual(
            (parsed["decor"], parsed["structure"], parsed["thickness"], parsed["length"]),
            (decor, structure, thickness, length),
            name,
        )

    def test_site_names(self) -> None:
        self.check("D/4230 SE/Джексон Гикори светлый 2.75*1.83", "4230", "SE", 16, 2750)
        self.check("D/U112 PE/Серый 2.75*1.83 v3", "U112", "PE", 16, 2750)
        self.check("D/20337 MX/Орех Белладжио 2,75*1,83", "20337", "MX", 16, 2750)
        self.check("D/K110 SM/Белый корпус 2.75*1.83", "K110", "SM", 16, 2750)
        self.check("D/D722 SE/Орех 2.75*1.83", "D722", "SE", 16, 2750)

    def test_cyrillic_lookalikes(self) -> None:
        self.check("D/4075 РЕ/Голубой 2.75*1.83", "4075", "PE", 16, 2750)
        self.check("D/1039/Орех Южный/ОW/2800*2070*18", "1039", "OW", 18, 2800)

    def test_factory_writes_some_codes_without_u(self) -> None:
        self.check("D/U6933 VL/Кашемир  2.75*1.83", "6933", "VL", 16, 2750)
        self.check("D/U10257 VL/Тоффи (Грейдж)  2.75*1.83", "10257", "VL", 16, 2750)
        # The factory's list has no 6931, so the site's code stays and the line gets flagged.
        self.check("D/U6931/Лемонграсс/VL/2800*2070*18", "U6931", "VL", 18, 2800)

    def test_other_goods_are_not_parsed(self) -> None:
        self.assertIsNone(krono.parse_product("Мебельная плита 8723/Bc Elain 2440х1830х18"))


class SupplierTest(unittest.TestCase):
    def test_names(self) -> None:
        for name in ("Swiss Krono", "SWISS KRONO", "Krono", "Свисс Кроно", "ОсОО Кроно"):
            self.assertTrue(krono.is_swiss_krono_supplier(name), name)
        for name in ("WQ", "Kronospan", "Кроношпан", "", None):
            self.assertFalse(krono.is_swiss_krono_supplier(name), name)


class BuildRowsTest(unittest.TestCase):
    def test_nearest_pallet(self) -> None:
        for sheets, pallets in ((0, 0), (16, 0), (16.5, 1), (17, 1), (33, 1), (49, 1), (50, 2)):
            rows = krono.build_order_rows(forecast(("D/4230 OW/Гикори 2.75*1.83", "Swiss Krono", sheets, "Высокий")))
            self.assertEqual(rows.iloc[0]["pallets"], pallets, sheets)

    def test_only_goods_of_the_supplier(self) -> None:
        rows = krono.build_order_rows(
            forecast(
                ("D/4230 OW/Гикори 2.75*1.83", "Swiss Krono", 33, "Высокий"),
                ("D/4080 OW/Другой 2.75*1.83", "WQ", 33, "Высокий"),
            )
        )
        self.assertEqual(rows["decor"].tolist(), ["4230"])

    def test_low_priority_is_not_proposed(self) -> None:
        rows = krono.build_order_rows(forecast(("D/4230 OW/Гикори 2.75*1.83", "Krono", 330, "Низкий")))
        self.assertEqual(rows.iloc[0]["pallets"], 0)
        self.assertTrue(rows.iloc[0]["low_priority_skipped"])

    def test_full_pallet_is_preferred_where_a_half_one_also_exists(self) -> None:
        rows = krono.build_order_rows(forecast(("D/4084 VL/Декор 2.75*1.83", "Krono", 33, "Высокий")))
        self.assertEqual(rows.iloc[0]["pack"], 33)
        self.assertEqual(rows.iloc[0]["index"], "D10164084VL2750001")

    def test_pack_of_the_factory_differs_from_the_form_formula(self) -> None:
        rows = krono.build_order_rows(forecast(("D/4041 SM/Декор 2750*1830*12", "Krono", 88, "Высокий")))
        self.assertEqual(rows.iloc[0]["pack"], 44)
        self.assertEqual(rows.iloc[0]["pallets"], 2)

    def test_unknown_material_is_flagged_and_left_out(self) -> None:
        rows = krono.build_order_rows(
            forecast(("D/99999 OW/Нет такого 2.75*1.83", "Krono", 330, "Высокий"), ("Просто товар", "Krono", 330, "Высокий"))
        )
        self.assertFalse(rows["in_catalog"].any())
        self.assertEqual(rows["pallets"].sum(), 0)
        self.assertTrue(krono.order_items_from_rows(rows).empty)

    def test_order_lines_are_in_sheets(self) -> None:
        rows = krono.build_order_rows(forecast(("D/4230 OW/Гикори 2.75*1.83", "Krono", 66, "Высокий")))
        items = krono.order_items_from_rows(rows)
        self.assertEqual(items.iloc[0]["product"], "D/4230 OW/Гикори 2.75*1.83")
        self.assertEqual(items.iloc[0]["quantity"], 66.0)


class TemplateUploadTest(unittest.TestCase):
    def setUp(self) -> None:
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        patch = mock.patch.object(krono, "TEMPLATE_PATH", Path(folder.name) / "forms" / "form.xlsx")
        patch.start()
        self.addCleanup(patch.stop)

    def test_factory_form_is_kept_and_used(self) -> None:
        self.assertFalse(krono.template_ready())
        krono.save_template(PRIVATE_TEMPLATE.read_bytes())
        self.assertTrue(krono.template_ready())
        rows = krono.build_order_rows(forecast(("D/4230 OW/Гикори 2.75*1.83", "Krono", 33, "Высокий")))
        self.assertEqual(rows.iloc[0]["pallets"], 1)
        self.assertEqual(workbook(krono.build_order_workbook(rows))[krono.ORDER_SHEET]["W155"].value, 33)
        self.assertEqual([path.name for path in krono.TEMPLATE_PATH.parent.iterdir()], ["form.xlsx"])

    def test_other_files_are_refused_and_change_nothing(self) -> None:
        krono.save_template(PRIVATE_TEMPLATE.read_bytes())
        before = krono.TEMPLATE_PATH.read_bytes()
        other = Workbook()
        other.active.append(["Номенклатура", "Количество"])
        buffer = BytesIO()
        other.save(buffer)
        for content in (b"not a workbook", buffer.getvalue()):
            with self.assertRaises(ValueError):
                krono.save_template(content)
        self.assertEqual(krono.TEMPLATE_PATH.read_bytes(), before)


class SavedOrderTest(unittest.TestCase):
    def test_saved_order_rebuilds_the_same_pallets(self) -> None:
        rows = krono.build_order_rows(
            forecast(
                ("D/4230 OW/Гикори 2.75*1.83", "Swiss Krono", 66, "Высокий"),
                ("D/K110 SM/Белый 2.75*1.83", "Swiss Krono", 396, "Высокий"),
            )
        )
        saved = krono.rows_from_saved_order(krono.order_items_from_rows(rows))
        self.assertEqual(dict(zip(saved["decor"], saved["pallets"])), {"K110": 12, "4230": 2})
        self.assertEqual(krono.order_totals(saved)["sheets"], 66 + 396)


class WorkbookTest(unittest.TestCase):
    ROWS = (
        ("D/4230 OW/Гикори 2.75*1.83", "Swiss Krono", 66, "Высокий"),
        ("D/K110 SM/Белый корпус 2.75*1.83", "Swiss Krono", 396, "Высокий"),
        ("D/U112 PE/Серый 2.75*1.83", "Swiss Krono", 33, "Высокий"),
        ("D/4041 SM/Особая паллета 2750*1830*12", "Swiss Krono", 88, "Высокий"),
        ("D/1036 VL/Не нужен 2.75*1.83", "Swiss Krono", 1, "Высокий"),
    )

    @classmethod
    def setUpClass(cls) -> None:
        cls.order_rows = krono.build_order_rows(forecast(*cls.ROWS))
        cls.content = krono.build_order_workbook(cls.order_rows)
        cls.values = workbook(cls.content)[krono.ORDER_SHEET]
        cls.formulas = workbook(cls.content, values=False)[krono.ORDER_SHEET]

    def rows_by_decor(self) -> dict[object, int]:
        return {
            self.values[f"E{row}"].value: row
            for row in range(krono.FIRST_ROW, krono.LAST_ROW + 1)
            if self.values[f"E{row}"].value is not None
        }

    def test_goods_are_written_with_what_the_form_expects(self) -> None:
        rows = self.rows_by_decor()
        self.assertEqual(set(rows), {4230, 110, 112, 4041})
        row = rows[110]
        self.assertEqual(self.values[f"G{row}"].value, "SM")
        self.assertEqual(self.values[f"I{row}"].value, 16)
        self.assertEqual(self.values[f"J{row}"].value, 2750)
        self.assertEqual(self.values[f"N{row}"].value, 12)
        self.assertEqual(self.values[f"W{row}"].value, 396)
        self.assertEqual(self.values[f"L{row}"].value, "D1016K110SM2750001")
        self.assertEqual(self.values[f"AA{row}"].value, "D/K110/K110/SM/SM//16,0/2750x1830/033/1")

    def test_stored_keys_exist_in_the_factory_list(self) -> None:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            listed = {str(row[0]).upper() for row in load_workbook(krono.TEMPLATE_PATH, data_only=True)["C"].iter_rows(values_only=True)}
        for row in self.rows_by_decor().values():
            self.assertIn(self.values[f"AA{row}"].value.upper(), listed)

    def test_factory_pack_is_typed_over_the_formula(self) -> None:
        row = self.rows_by_decor()[4041]
        self.assertEqual(self.formulas[f"X{row}"].value, 44)
        self.assertEqual(self.values[f"W{row}"].value, 88)
        self.assertEqual(self.values[f"AA{row}"].value, "D/4041/4041/SM/SM//12,0/2750x1830/044/1")

    def test_totals_ignore_the_number_the_blank_had_typed_in(self) -> None:
        self.assertEqual(self.values["W155"].value, 66 + 396 + 33 + 88)
        self.assertEqual(self.values["N155"].value, 2 + 12 + 1 + 2)
        self.assertAlmostEqual(self.values["Y155"].value, (66 + 396 + 33 + 88) * 5.0325)
        self.assertTrue(str(self.formulas["W21"].value).startswith("=M21+O21"))

    def test_the_sample_order_of_the_blank_is_gone(self) -> None:
        self.assertNotIn(20337, self.rows_by_decor())
        self.assertNotIn(4080, self.rows_by_decor())

    def test_formulas_lists_and_protection_of_the_blank_stay(self) -> None:
        self.assertTrue(str(self.formulas["M5"].value).startswith("=X5*N5"))
        self.assertTrue(str(self.formulas["L5"].value).startswith("=VLOOKUP"))
        self.assertTrue(self.formulas.protection.sheet)
        with ZipFile(BytesIO(self.content)) as archive:
            names = set(archive.namelist())
            sheet = archive.read(krono.ORDER_SHEET_PART).decode("utf-8")
            types = archive.read("[Content_Types].xml").decode("utf-8")
            relations = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            book = archive.read("xl/workbook.xml").decode("utf-8")
        self.assertIn("x14:conditionalFormatting", sheet)
        self.assertIn("x14:dataValidation", sheet)
        self.assertNotIn("xl/calcChain.xml", names)
        self.assertNotIn("calcChain", types + relations)
        self.assertIn('fullCalcOnLoad="1"', book)
        self.assertEqual(workbook(self.content)["C"].sheet_state, "hidden")

    def test_shares_match_the_price_groups(self) -> None:
        totals = krono.order_totals(self.order_rows)
        self.assertAlmostEqual(self.values["N1"].value, totals["middle_share"])
        self.assertAlmostEqual(self.values["P1"].value, totals["top_share"])
        _, middle, top = krono.load_catalog()
        self.assertEqual((middle, top), (0.2, 0.05))

    def test_empty_order_still_gives_a_clean_blank(self) -> None:
        rows = krono.build_order_rows(forecast(("D/1036 VL/Не нужен 2.75*1.83", "Krono", 0, "Высокий")))
        values = workbook(krono.build_order_workbook(rows))[krono.ORDER_SHEET]
        self.assertEqual(values["W155"].value, 0)


class FileRoundTripTest(unittest.TestCase):
    def setUp(self) -> None:
        self.rows = krono.build_order_rows(
            forecast(
                ("D/4230 OW/Гикори 2.75*1.83", "Swiss Krono", 66, "Высокий"),
                ("D/K110 SM/Белый 2.75*1.83", "Swiss Krono", 396, "Высокий"),
                ("D/4080 OW/Предложен ноль 2.75*1.83", "Swiss Krono", 0, "Высокий"),
            )
        )
        self.content = krono.build_order_workbook(self.rows)

    def edited(self, change) -> bytes:
        book = workbook(self.content, values=False)
        change(book[krono.ORDER_SHEET])
        output = BytesIO()
        book.save(output)
        return output.getvalue()

    def by_decor(self, frame: pd.DataFrame) -> dict[str, float]:
        return dict(zip(frame["decor"], frame["pallets"]))

    def test_unchanged_file_gives_the_same_order(self) -> None:
        parsed = krono.parse_order_file(self.content, self.rows)
        self.assertEqual((parsed.errors, parsed.unknown), ([], []))
        self.assertEqual(self.by_decor(krono.apply_order_file(self.rows, parsed)), self.by_decor(self.rows))

    def test_changes_second_half_new_line_and_removal(self) -> None:
        def change(sheet) -> None:
            for row in range(krono.FIRST_ROW, krono.LAST_ROW):
                if sheet[f"E{row}"].value == 110:
                    sheet[f"N{row}"] = 5
                    sheet[f"P{row}"] = 2
                if sheet[f"E{row}"].value == 4230:
                    sheet[f"N{row}"] = None
            sheet["E40"], sheet["G40"], sheet["I40"], sheet["J40"], sheet["N40"] = 4080, "OW", 16, 2750, 3

        result = krono.apply_order_file(self.rows, krono.parse_order_file(self.edited(change), self.rows))
        self.assertEqual(self.by_decor(result), {"K110": 7, "4230": 0, "4080": 3})

    def test_line_not_among_the_goods_is_reported(self) -> None:
        def change(sheet) -> None:
            sheet["E40"], sheet["G40"], sheet["I40"], sheet["J40"], sheet["N40"] = 9999, "OW", 16, 2750, 1

        parsed = krono.parse_order_file(self.edited(change), self.rows)
        self.assertEqual(parsed.unknown, ["D/9999 OW 16 мм"])

    def test_bad_number_and_wrong_file(self) -> None:
        def change(sheet) -> None:
            for row in range(krono.FIRST_ROW, krono.LAST_ROW):
                if sheet[f"E{row}"].value == 110:
                    sheet[f"N{row}"] = "много"

        self.assertTrue(krono.parse_order_file(self.edited(change), self.rows).errors)
        self.assertTrue(krono.parse_order_file(b"not a workbook", self.rows).errors)


if __name__ == "__main__":
    unittest.main()
