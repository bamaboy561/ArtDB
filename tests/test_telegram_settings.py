from __future__ import annotations

import os
from datetime import date
from io import BytesIO
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from urllib.parse import parse_qs

import pandas as pd
from openpyxl import load_workbook
from PIL import Image


APP_DIR = Path(__file__).resolve().parents[1] / "app"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

from telegram_reports import (
    TelegramReportFile,
    build_targeted_telegram_card,
    build_targeted_telegram_report,
    discover_telegram_chats,
    send_telegram_document,
    send_telegram_message,
    send_telegram_photo,
)
from telegram_settings_store import (
    add_telegram_chat_id,
    load_environment_telegram_settings,
    parse_telegram_chat_ids,
    remove_telegram_chat_id,
)


class TelegramEnvironmentSettingsTests(unittest.TestCase):
    def test_environment_settings_are_normalized(self) -> None:
        environment = {
            "TG_BOT_TOKEN": " token ",
            "TG_CHAT_ID": " -100123 ",
            "TELEGRAM_DAILY_ENABLED": "true",
            "TELEGRAM_DAILY_REPORT_HOUR": "25",
            "TELEGRAM_DAILY_REPORT_MINUTE": "-4",
            "TELEGRAM_SEND_REPORT_FILES": "yes",
        }
        with patch.dict(os.environ, environment, clear=True):
            settings = load_environment_telegram_settings()

        self.assertEqual(settings.bot_token, "token")
        self.assertEqual(settings.chat_id, "-100123")
        self.assertTrue(settings.configured)
        self.assertTrue(settings.daily_enabled)
        self.assertEqual(settings.report_hour, 23)
        self.assertEqual(settings.report_minute, 0)
        self.assertTrue(settings.send_report_files)

    def test_multiple_chat_ids_are_normalized_and_deduplicated(self) -> None:
        environment = {
            "TG_BOT_TOKEN": "token",
            "TG_CHAT_IDS": "123, -100555\n123;456",
        }
        with patch.dict(os.environ, environment, clear=True):
            settings = load_environment_telegram_settings()

        self.assertEqual(settings.chat_ids, ("123", "-100555", "456"))
        self.assertEqual(settings.chat_id, "123,-100555,456")

    def test_chat_list_helpers_append_and_remove_without_duplicates(self) -> None:
        connected = add_telegram_chat_id("123,-100555", "123")
        self.assertEqual(parse_telegram_chat_ids(connected), ("123", "-100555"))
        self.assertEqual(remove_telegram_chat_id(connected, "123"), "-100555")


class TelegramChatDiscoveryTests(unittest.TestCase):
    def test_discovery_returns_unique_recent_chats(self) -> None:
        response = {
            "ok": True,
            "result": [
                {
                    "update_id": 1,
                    "message": {
                        "chat": {
                            "id": 123,
                            "type": "private",
                            "first_name": "Sultan",
                            "username": "sultan",
                        }
                    },
                },
                {
                    "update_id": 2,
                    "message": {
                        "chat": {
                            "id": 123,
                            "type": "private",
                            "first_name": "Sultan",
                        }
                    },
                },
                {
                    "update_id": 3,
                    "message": {
                        "chat": {
                            "id": -100555,
                            "type": "supergroup",
                            "title": "ArtDB reports",
                        }
                    },
                },
            ],
        }

        with patch("telegram_reports._telegram_api_request", return_value=response):
            chats = discover_telegram_chats()

        self.assertEqual({item["chat_id"] for item in chats}, {"123", "-100555"})
        labels = {item["chat_id"]: item["label"] for item in chats}
        self.assertEqual(labels["123"], "Sultan")
        self.assertEqual(labels["-100555"], "ArtDB reports")


class TelegramMessageSafetyTests(unittest.TestCase):
    def test_message_escapes_unsupported_html_without_losing_bold_text(self) -> None:
        with (
            patch("telegram_reports._get_telegram_credentials", return_value=("token", "123")),
            patch("telegram_reports._telegram_api_request", return_value={"ok": True}) as api_request,
        ):
            send_telegram_message("<b>Риски</b> (<15%): A&B; товар &lt;b&gt;SKU&lt;/b&gt;")

        payload = api_request.call_args.args[1]
        fields = parse_qs(payload.decode("utf-8"))
        self.assertEqual(
            fields["text"],
            ["<b>Риски</b> (&lt;15%): A&amp;B; товар &lt;b&gt;SKU&lt;/b&gt;"],
        )
        self.assertEqual(fields["parse_mode"], ["HTML"])

    def test_message_supports_target_chat_and_keyboard(self) -> None:
        keyboard = {"inline_keyboard": [[{"text": "30 дней", "callback_data": "period:30"}]]}
        with (
            patch("telegram_reports._get_telegram_credentials", return_value=("token", "123")),
            patch("telegram_reports._telegram_api_request", return_value={"ok": True}) as api_request,
        ):
            send_telegram_message("Отчёт", chat_id="456", reply_markup=keyboard)

        fields = parse_qs(api_request.call_args.args[1].decode("utf-8"))
        self.assertEqual(fields["chat_id"], ["456"])
        self.assertIn('"callback_data": "period:30"', fields["reply_markup"][0])

    def test_message_is_broadcast_to_all_configured_chats(self) -> None:
        with (
            patch("telegram_reports._get_telegram_credentials", return_value=("token", "123,-100555")),
            patch("telegram_reports._telegram_api_request", return_value={"ok": True}) as api_request,
        ):
            send_telegram_message("Отчёт")

        target_ids = [
            parse_qs(call.args[1].decode("utf-8"))["chat_id"][0]
            for call in api_request.call_args_list
        ]
        self.assertEqual(target_ids, ["123", "-100555"])

    def test_document_is_broadcast_to_all_configured_chats(self) -> None:
        report_file = TelegramReportFile(
            filename="report.xlsx",
            content=b"report-content",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        with (
            patch("telegram_reports._get_telegram_credentials", return_value=("token", "123,-100555")),
            patch("telegram_reports._telegram_api_request", return_value={"ok": True}) as api_request,
        ):
            send_telegram_document(report_file)

        self.assertEqual(api_request.call_count, 2)
        payloads = [call.args[1] for call in api_request.call_args_list]
        self.assertIn(b'\r\n\r\n123\r\n', payloads[0])
        self.assertIn(b'\r\n\r\n-100555\r\n', payloads[1])

    def test_photo_is_broadcast_to_all_configured_chats(self) -> None:
        report_file = TelegramReportFile(
            filename="summary.png",
            content=b"png-content",
            caption="<b>ArtDB</b>",
            content_type="image/png",
        )
        with (
            patch("telegram_reports._get_telegram_credentials", return_value=("token", "123,-100555")),
            patch("telegram_reports._telegram_api_request", return_value={"ok": True}) as api_request,
        ):
            send_telegram_photo(report_file)

        self.assertEqual(api_request.call_count, 2)
        self.assertTrue(all(call.args[0] == "sendPhoto" for call in api_request.call_args_list))
        self.assertTrue(all(b'name="photo"' in call.args[1] for call in api_request.call_args_list))


class TargetedTelegramReportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.sales = pd.DataFrame(
            [
                {
                    "date": pd.Timestamp("2026-06-15"),
                    "salon": "Artisan",
                    "item_code": "A-1",
                    "product_key": "A-1",
                    "product": "Плита Дуб",
                    "category": "ЛДСП",
                    "brand": "Egger",
                    "supplier": "Slotex",
                    "manager": "Айбек",
                    "quantity": 1.0,
                    "revenue": 100.0,
                    "cost": 60.0,
                    "margin": 40.0,
                    "margin_pct": 40.0,
                },
                {
                    "date": pd.Timestamp("2026-07-05"),
                    "salon": "Artisan",
                    "item_code": "A-1",
                    "product_key": "A-1",
                    "product": "Плита Дуб",
                    "category": "ЛДСП",
                    "brand": "Egger",
                    "supplier": "Slotex",
                    "manager": "Айбек",
                    "quantity": 2.0,
                    "revenue": 300.0,
                    "cost": 180.0,
                    "margin": 120.0,
                    "margin_pct": 40.0,
                },
                {
                    "date": pd.Timestamp("2026-07-06"),
                    "salon": "Artisan",
                    "item_code": "B-2",
                    "product_key": "B-2",
                    "product": "Петля <15 мм",
                    "category": "Фурнитура",
                    "brand": "Hettich",
                    "supplier": "Hettich",
                    "manager": "Бек",
                    "quantity": 4.0,
                    "revenue": 200.0,
                    "cost": 140.0,
                    "margin": 60.0,
                    "margin_pct": 30.0,
                },
            ]
        )

    def test_category_report_respects_period_and_category(self) -> None:
        message, report_file = build_targeted_telegram_report(
            self.sales,
            report_kind="portfolio",
            date_from=date(2026, 7, 1),
            date_to=date(2026, 7, 31),
            category="ЛДСП",
        )

        self.assertIn("Портфель SKU", message)
        self.assertIn("Категория: ЛДСП", message)
        self.assertIn("300 сом", message)
        self.assertIn("━━━━━━━━━━━━", message)
        self.assertIn("<b>Ключевые показатели</b>", message)
        self.assertIn("<b>Сравнение с предыдущим периодом</b>", message)
        self.assertIn("<b>1. A-1", message)
        self.assertIn("</b>\n   Выручка:", message)
        workbook = load_workbook(BytesIO(report_file.content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames[:3], ["Сводка", "Портфель SKU", "Динамика"])
        portfolio_rows = list(workbook["Портфель SKU"].iter_rows(values_only=True))
        self.assertEqual(len(portfolio_rows), 2)
        self.assertIn("A-1", portfolio_rows[1])
        workbook.close()

        styled_workbook = load_workbook(BytesIO(report_file.content), read_only=False, data_only=True)
        portfolio_sheet = styled_workbook["Портфель SKU"]
        self.assertEqual(portfolio_sheet.freeze_panes, "A2")
        self.assertEqual(portfolio_sheet["A1"].fill.fgColor.rgb, "00003461")
        self.assertTrue(bool(portfolio_sheet.auto_filter.ref))
        styled_workbook.close()

    def test_targeted_visual_card_is_a_readable_png(self) -> None:
        report_file = build_targeted_telegram_card(
            self.sales,
            report_kind="summary",
            date_from=date(2026, 6, 1),
            date_to=date(2026, 7, 31),
        )

        self.assertEqual(report_file.content_type, "image/png")
        self.assertTrue(report_file.content.startswith(b"\x89PNG\r\n\x1a\n"))
        with Image.open(BytesIO(report_file.content)) as image:
            self.assertEqual(image.size, (1200, 830))
            self.assertEqual(image.mode, "RGB")

    def test_sku_report_contains_only_selected_article_sales(self) -> None:
        message, report_file = build_targeted_telegram_report(
            self.sales,
            report_kind="sku",
            date_from=date(2026, 7, 1),
            date_to=date(2026, 7, 31),
            product_key="B-2",
        )

        self.assertIn("B-2", message)
        workbook = load_workbook(BytesIO(report_file.content), read_only=True, data_only=True)
        detail_rows = list(workbook["Продажи SKU"].iter_rows(values_only=True))
        self.assertEqual(len(detail_rows), 2)
        self.assertIn("B-2", detail_rows[1])
        self.assertNotIn("A-1", detail_rows[1])
        workbook.close()

    def test_message_does_not_promise_excel_when_file_is_disabled(self) -> None:
        message, _ = build_targeted_telegram_report(
            self.sales,
            report_kind="summary",
            date_from=date(2026, 7, 1),
            date_to=date(2026, 7, 31),
            include_file_note=False,
        )

        self.assertNotIn("Excel-файле", message)

    def test_supplier_report_includes_inventory_and_reorder_sheets(self) -> None:
        forecast = pd.DataFrame(
            [
                {
                    "product": "Плита Дуб",
                    "category": "ЛДСП",
                    "supplier": "Slotex",
                    "brand": "Egger",
                    "abc_class": "A",
                    "xyz_class": "X",
                    "priority": "Критичный",
                    "stock_status": "Риск дефицита",
                    "forecast_qty": 10.0,
                    "stock_on_hand": 2.0,
                    "stock_value": 500.0,
                    "manual_stock_in_transit": 0.0,
                    "ordered_in_transit_qty": 1.0,
                    "stock_in_transit": 1.0,
                    "available_stock_qty": 3.0,
                    "stock_coverage_days": 12.0,
                    "coverage_requirement_qty": 10.0,
                    "gross_requirement_qty": 12.0,
                    "net_requirement_qty": 9.0,
                    "recommended_order_qty": 9.0,
                    "lead_time_days": 14,
                    "last_sale_date": pd.Timestamp("2026-07-05"),
                    "days_since_last_sale": 2,
                    "forecast_revenue": 1500.0,
                }
            ]
        )

        message, report_file = build_targeted_telegram_report(
            self.sales,
            report_kind="supplier",
            date_from=date(2026, 7, 1),
            date_to=date(2026, 7, 31),
            supplier="Slotex",
            procurement_forecast=forecast,
        )

        self.assertIn("Поставщик: Slotex", message)
        self.assertIn("500 сом", message)
        self.assertIn("<b>Остаток, шт.:</b> 2", message)
        self.assertIn("<b>Стоимость остатка:</b> 500 сом", message)
        self.assertIn("<b>Дефицит:</b> 1 SKU", message)
        workbook = load_workbook(BytesIO(report_file.content), read_only=True, data_only=True)
        self.assertIn("Остатки и прогноз", workbook.sheetnames)
        self.assertIn("К заказу", workbook.sheetnames)
        self.assertEqual(len(list(workbook["К заказу"].iter_rows(values_only=True))), 2)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
