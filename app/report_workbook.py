"""Readable Excel workbooks for reports that people open on a phone or a laptop.

Every data sheet gets a title, a one-line explanation, a styled header, number
and date formats, zebra rows and a totals row. A key/value "Сводка" sheet is
rendered as a cover page with the key figures and a table of contents.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import math
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

NAVY = "003461"
GREEN = "006C49"
TEXT = "1E293B"
MUTED = "64748B"
ZEBRA = "F5F8FB"
TOTAL_FILL = "E6EEF6"
LINE = "D8E0E8"
RED = "B91C1C"
POSITIVE = "047857"

HEADER_ROW = 4
FIRST_DATA_ROW = HEADER_ROW + 1
COVER_COLUMNS = ["Показатель", "Значение", "Единица"]
RU_MONTHS = (
    "январь", "февраль", "март", "апрель", "май", "июнь",
    "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
)
MONTH_KEY = re.compile(r"^(20\d{2})-(0[1-9]|1[0-2])$")
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")

MONEY_MARKERS = ("выруч", "себесто", "прибыл", "марж", "сумм", "стоимост", "цена")
NO_TOTAL_MARKERS = ("%", "дней", "дн.", "срок", "цена", "покрыт", "тренд", "средн", "кратн", "moq", "вариатив")
TOTAL_MARKERS = (
    "выруч", "себесто", "прибыл", "марж", "сумм", "стоимост", "колич", "кол-во", "шт",
    "заказ", "потребн", "остат", "доступ", "в пути", "sku", "строк", "прогноз",
)
STATUS_FILLS = (
    (("критич", "нет остатка", "дефицит", "просроч"), "FDE2E4", "991B1B"),
    (("высок", "пополнен", "излиш", "неликвид", "нет спроса"), "FEF3C7", "92400E"),
    (("покрыт", "норма", "получен"), "DCFCE7", "166534"),
)

SHEET_DESCRIPTIONS = {
    "динамика": "Продажи по месяцам: выручка, прибыль и количество, изменение к предыдущему месяцу.",
    "по месяцам": "Продажи по месяцам: выручка, прибыль и количество, изменение к предыдущему месяцу.",
    "динамика sku": "Продажи выбранного товара по месяцам.",
    "категории": "Вклад каждой категории в выручку и прибыль за период.",
    "портфель sku": "Все товары периода, отсортированные по выручке: сколько продано и сколько заработано.",
    "топ товаров": "Товары с наибольшей выручкой за всё время в архиве.",
    "карточка sku": "Итог по выбранному товару за период.",
    "продажи sku": "Каждая строка продаж выбранного товара: дата, салон, менеджер, сумма.",
    "продажи": "Каждая строка продаж за период: дата, салон, товар, менеджер, сумма.",
    "остатки и прогноз": "Остаток, прогноз спроса и рекомендация к заказу по каждому товару.",
    "прогноз закупок": "Остаток, прогноз спроса и рекомендация к заказу по каждому товару.",
    "к заказу": "Товары, которые уже пора заказывать, и рекомендуемое количество.",
    "заказ общий": "Все позиции к заказу по всем поставщикам одним списком.",
    "дефицит": "Товары, которых не хватит до следующей поставки.",
    "излишки": "Товары, которых на складе больше, чем нужно на ближайший период.",
    "остатки и риски": "Проблемные позиции: дефицит, излишки, неликвид и товары к заказу.",
    "поставщики": "Сколько позиций и единиц нужно заказать у каждого поставщика.",
    "по поставщикам": "Сколько позиций и единиц нужно заказать у каждого поставщика.",
    "загрузки": "Какие файлы продаж загружены в архив: салон, дата отчёта, имя файла.",
}

THIN = Side(style="thin", color=LINE)
HEADER_SIDE = Side(style="thin", color="FFFFFF")


def humanize_month(value: object) -> object:
    match = MONTH_KEY.match(str(value).strip())
    if not match:
        return value
    return f"{RU_MONTHS[int(match.group(2)) - 1]} {match.group(1)}"


def describe_sheet(sheet_name: str, descriptions: dict[str, str] | None = None) -> str:
    if descriptions and sheet_name in descriptions:
        return descriptions[sheet_name]
    return SHEET_DESCRIPTIONS.get(sheet_name.casefold().strip(), "")


def safe_sheet_name(value: object, used_names: set[str]) -> str:
    raw_name = str(value or "Лист").strip() or "Лист"
    safe_name = re.sub(r"[\[\]\:\*\?\/\\]", " ", raw_name)
    safe_name = re.sub(r"\s+", " ", safe_name).strip()[:31] or "Лист"
    candidate = safe_name
    suffix = 2
    while candidate.casefold() in used_names:
        suffix_text = f" {suffix}"
        candidate = f"{safe_name[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used_names.add(candidate.casefold())
    return candidate


def _clean_value(value: object, header: str = "") -> object:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if value is pd.NaT or value is pd.NA:
        return None
    if isinstance(value, pd.Timestamp):
        return None if pd.isna(value) else value.to_pydatetime().date()
    if isinstance(value, datetime):
        return value.date() if (value.hour, value.minute, value.second) == (0, 0, 0) else value
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        try:
            value = value.item()
        except (ValueError, AttributeError):
            pass
        if isinstance(value, float) and math.isnan(value):
            return None
    if isinstance(value, str):
        text = value.strip()
        if MONTH_KEY.match(text):
            return humanize_month(text)
        if ISO_DATE.match(text) and _is_date_header(header):
            parsed = pd.to_datetime(text, errors="coerce")
            if pd.notna(parsed):
                return parsed.to_pydatetime() if "загруж" in header.casefold() else parsed.date()
        return text
    return value


def _is_date_header(header: str) -> bool:
    lowered = header.casefold()
    return any(marker in lowered for marker in ("дата", "продажа", "загруж", "период", "обновл"))


def _is_money_header(header: str) -> bool:
    lowered = header.casefold()
    return "%" not in lowered and any(marker in lowered for marker in MONEY_MARKERS)


def _is_total_header(header: str) -> bool:
    lowered = header.casefold()
    if any(marker in lowered for marker in NO_TOTAL_MARKERS):
        return False
    return any(marker in lowered for marker in TOTAL_MARKERS)


def _number_format(header: str, values: list[object]) -> str:
    if "%" in header:
        return '0.0"%"'
    numbers = [value for value in values if isinstance(value, (int, float)) and not isinstance(value, bool)]
    whole = all(float(value).is_integer() for value in numbers) if numbers else True
    if _is_money_header(header):
        return "#,##0" if whole or max((abs(float(v)) for v in numbers), default=0) >= 1000 else "#,##0.00"
    return "#,##0" if whole else "#,##0.0"


def _write_title(sheet: Worksheet, title: str, subtitle: str, width: int) -> None:
    last_column = max(width, 4)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    title_cell = sheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    title_cell.alignment = Alignment(vertical="center")
    subtitle_cell = sheet.cell(row=2, column=1, value=subtitle)
    subtitle_cell.font = Font(name="Calibri", size=10, color=MUTED)
    subtitle_cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 30 if len(subtitle) > 90 else 18
    sheet.row_dimensions[3].height = 6


def _setup_print(sheet: Worksheet, *, landscape: bool, repeat_header: bool) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4
    sheet.page_margins.top = sheet.page_margins.bottom = 0.5
    if repeat_header:
        sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    sheet.oddFooter.center.text = "Стр. &P из &N"
    sheet.oddFooter.center.size = 8


def write_table_sheet(sheet: Worksheet, frame: pd.DataFrame, *, title: str, subtitle: str) -> None:
    headers = [str(column) for column in frame.columns]
    column_count = max(len(headers), 1)
    _write_title(sheet, title, subtitle, column_count)
    sheet.sheet_properties.tabColor = NAVY

    if not headers or frame.empty:
        note = sheet.cell(row=HEADER_ROW, column=1, value="За выбранный период данных для этого листа нет.")
        note.font = Font(name="Calibri", size=11, italic=True, color=MUTED)
        sheet.column_dimensions["A"].width = 60
        _setup_print(sheet, landscape=False, repeat_header=False)
        return

    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=column_index, value=header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=HEADER_SIDE, right=HEADER_SIDE)
    sheet.row_dimensions[HEADER_ROW].height = 32

    columns_values: list[list[object]] = []
    for column_index, header in enumerate(headers, start=1):
        columns_values.append([_clean_value(value, header) for value in frame.iloc[:, column_index - 1].tolist()])

    row_count = len(frame)
    last_data_row = FIRST_DATA_ROW + row_count - 1
    for column_index, (header, values) in enumerate(zip(headers, columns_values), start=1):
        number_format = _number_format(header, values)
        is_change = "изменен" in header.casefold() or "откл" in header.casefold()
        longest = len(header) + 3
        texts = [value for value in values if isinstance(value, str) and value]
        is_code_column = bool(texts) and all(len(text) <= 3 for text in texts)
        for offset, value in enumerate(values):
            row_index = FIRST_DATA_ROW + offset
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.border = Border(bottom=THIN)
            font_color = TEXT
            if offset % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=ZEBRA)
            if isinstance(value, bool):
                cell.value = "да" if value else "нет"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif isinstance(value, (int, float)):
                cell.number_format = number_format
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if value < 0:
                    font_color = RED
                elif is_change and value > 0:
                    font_color = POSITIVE
                longest = max(longest, len(f"{value:,.0f}") + 3)
            elif isinstance(value, datetime):
                cell.number_format = "DD.MM.YYYY HH:MM"
                cell.alignment = Alignment(horizontal="center", vertical="center")
                longest = max(longest, 16)
            elif isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
                cell.alignment = Alignment(horizontal="center", vertical="center")
                longest = max(longest, 11)
            else:
                text = "" if value is None else str(value)
                cell.alignment = Alignment(
                    horizontal="center" if is_code_column else None,
                    vertical="center",
                    wrap_text=len(text) > 48,
                )
                lowered = text.casefold()
                if text and len(text) <= 40:
                    for markers, fill_color, text_color in STATUS_FILLS:
                        if any(marker in lowered for marker in markers):
                            cell.fill = PatternFill("solid", fgColor=fill_color)
                            font_color = text_color
                            break
                if offset < 200:
                    longest = max(longest, min(len(text), 48))
            cell.font = Font(name="Calibri", size=10, color=font_color, bold=column_index == 1)
        minimum_width = 28 if column_index == 1 else 10
        sheet.column_dimensions[get_column_letter(column_index)].width = max(minimum_width, min(longest + 3, 52))

    if row_count >= 2:
        total_row = last_data_row + 1
        for column_index, (header, values) in enumerate(zip(headers, columns_values), start=1):
            cell = sheet.cell(row=total_row, column=column_index)
            cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
            cell.border = Border(top=Side(style="medium", color=NAVY))
            cell.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
            numbers = [value for value in values if isinstance(value, (int, float)) and not isinstance(value, bool)]
            if column_index == 1:
                cell.value = f"Итого ({row_count} строк)"
            elif numbers and _is_total_header(header):
                cell.value = float(sum(numbers))
                cell.number_format = _number_format(header, numbers + [cell.value])
                cell.alignment = Alignment(horizontal="right", vertical="center")
        sheet.row_dimensions[total_row].height = 20

    last_letter = get_column_letter(column_count)
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_letter}{last_data_row}"
    sheet.freeze_panes = f"B{FIRST_DATA_ROW}"
    _setup_print(sheet, landscape=column_count > 5, repeat_header=True)


def _cover_value_format(unit: str, value: object) -> str:
    unit = unit.strip()
    if unit == "%":
        return '0.0"%"'
    if unit in {"сом", "шт."}:
        return f'#,##0" {unit}"'
    if isinstance(value, float) and not float(value).is_integer():
        return "#,##0.0"
    return "#,##0"


def write_cover_sheet(
    sheet: Worksheet,
    frame: pd.DataFrame,
    *,
    title: str,
    subtitle: str,
    contents: list[tuple[str, str, int]],
) -> None:
    sheet.sheet_properties.tabColor = GREEN
    _write_title(sheet, title, subtitle, 3)
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 64

    parameters: list[tuple[str, object]] = []
    figures: list[tuple[str, object, str]] = []
    for label, value, unit in frame[COVER_COLUMNS].itertuples(index=False, name=None):
        clean = _clean_value(value, "дата" if "период" in str(label).casefold() else "")
        unit_text = "" if unit is None or (isinstance(unit, float) and math.isnan(unit)) else str(unit)
        if isinstance(clean, (int, float)) and not isinstance(clean, bool):
            # A zero stock value means the stock file was never loaded, not an empty warehouse.
            if clean == 0 and "стоимость" in str(label).casefold():
                clean = None
            figures.append((str(label), clean, unit_text))
        elif clean is None and unit_text:
            figures.append((str(label), None, unit_text))
        else:
            parameters.append((str(label), clean))

    row = HEADER_ROW

    def section(caption: str) -> None:
        nonlocal row
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        cell = sheet.cell(row=row, column=1, value=caption)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", indent=1)
        sheet.row_dimensions[row].height = 22
        row += 1

    def label_cell(text: str) -> None:
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font = Font(name="Calibri", size=10, color=MUTED)
        cell.alignment = Alignment(vertical="center", indent=1)
        cell.border = Border(bottom=THIN)

    if parameters:
        section("Параметры отчёта")
        for label, value in parameters:
            label_cell(label)
            cell = sheet.cell(row=row, column=2, value="—" if value in (None, "") else value)
            cell.font = Font(name="Calibri", size=11, bold=True, color=TEXT)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=THIN)
            if isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
            sheet.cell(row=row, column=3).border = Border(bottom=THIN)
            row += 1
        row += 1

    if figures:
        section("Ключевые показатели")
        for label, value, unit in figures:
            label_cell(label)
            cell = sheet.cell(row=row, column=2, value="нет данных" if value is None else value)
            is_change = "изменен" in label.casefold()
            color = NAVY
            if isinstance(value, (int, float)):
                cell.number_format = _cover_value_format(unit, value)
                if value < 0:
                    color = RED
                elif is_change and value > 0:
                    color = POSITIVE
            cell.font = Font(name="Calibri", size=13, bold=True, color=color)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = Border(bottom=THIN)
            note = sheet.cell(row=row, column=3)
            note.border = Border(bottom=THIN)
            if value is None:
                cell.font = Font(name="Calibri", size=11, color=MUTED)
            if is_change:
                note.value = (
                    "по сравнению с предыдущим периодом"
                    if value is not None
                    else "за предыдущий период продаж нет, сравнить не с чем"
                )
                note.font = Font(name="Calibri", size=9, italic=True, color=MUTED)
            sheet.row_dimensions[row].height = 21
            row += 1
        row += 1

    if contents:
        section("Что в этом файле")
        for sheet_name, description, rows in contents:
            link = sheet.cell(row=row, column=1, value=sheet_name)
            link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{sheet_name}'!A1", display=sheet_name)
            link.font = Font(name="Calibri", size=11, bold=True, color="1D6FC9", underline="single")
            link.alignment = Alignment(vertical="center", indent=1)
            link.border = Border(bottom=THIN)
            count = sheet.cell(row=row, column=2, value=f"{rows} строк" if rows else "нет данных")
            count.font = Font(name="Calibri", size=10, color=MUTED)
            count.alignment = Alignment(vertical="center")
            count.border = Border(bottom=THIN)
            about = sheet.cell(row=row, column=3, value=description)
            about.font = Font(name="Calibri", size=10, color=TEXT)
            about.alignment = Alignment(vertical="center", wrap_text=True)
            about.border = Border(bottom=THIN)
            sheet.row_dimensions[row].height = 30 if len(description) > 62 else 20
            row += 1
        row += 1
        hint = sheet.cell(row=row, column=1, value="Нажмите на название листа, чтобы перейти к нему.")
        hint.font = Font(name="Calibri", size=9, italic=True, color=MUTED)

    _setup_print(sheet, landscape=False, repeat_header=False)


def is_cover_frame(frame: pd.DataFrame) -> bool:
    return list(frame.columns) == COVER_COLUMNS


def build_report_workbook(
    sheets: dict[str, pd.DataFrame],
    *,
    title: str = "Отчёт ArtDB",
    subtitle: str = "",
    descriptions: dict[str, str] | None = None,
    generated_at: datetime | None = None,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    stamp = (generated_at or datetime.now()).strftime("%d.%m.%Y %H:%M")
    cover_subtitle = " · ".join(part for part in (subtitle, f"Сформирован {stamp}") if part)

    named = [(safe_sheet_name(raw_name, used_names), raw_name, frame) for raw_name, frame in sheets.items()]
    contents = [
        (sheet_name, describe_sheet(str(raw_name), descriptions), len(frame))
        for sheet_name, raw_name, frame in named
        if not is_cover_frame(frame)
    ]
    for sheet_name, raw_name, frame in named:
        sheet = workbook.create_sheet(sheet_name)
        if is_cover_frame(frame):
            write_cover_sheet(sheet, frame, title=title, subtitle=cover_subtitle, contents=contents)
        else:
            description = describe_sheet(str(raw_name), descriptions)
            write_table_sheet(
                sheet,
                frame,
                title=sheet_name,
                subtitle=" ".join(part for part in (description, f"({subtitle})" if subtitle else "") if part),
            )
    if not workbook.sheetnames:
        workbook.create_sheet("Отчёт")
    workbook.active = 0
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
