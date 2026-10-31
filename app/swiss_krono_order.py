"""The Swiss Krono order: the factory's own form, filled from the procurement forecast.

The form holds drop-down lists, conditional colours and sheet protection that
openpyxl would lose on saving, so its worksheet XML is edited directly. It belongs
to the factory and changes with its price lists, so it is not kept in the repository:
a manager uploads it on the site and it stays in the data folder.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from functools import lru_cache
import hashlib
from io import BytesIO
import math
import os
from pathlib import Path
import re
import warnings
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from openpyxl import load_workbook

from supplier_order_forms import EXCLUDED_ORDER_PRIORITY, _contains_term, _normalize


DATA_DIR = Path(os.getenv("APP_DATA_DIR", str(Path(__file__).resolve().parent.parent / "data"))).resolve()
TEMPLATE_PATH = DATA_DIR / "order_templates" / "swiss_krono_ldsp.xlsx"
ORDER_SHEET = "Заказ на ЛДСП"
ORDER_SHEET_PART = "xl/worksheets/sheet1.xml"
CATALOG_SHEET = "C"
SETTINGS_SHEET = "Справочник"
FIRST_ROW = 5
LAST_ROW = 154
TOTAL_ROW = 155
# The form multiplies sheets by this area for every format.
SHEET_AREA_M2 = 5.0325
PALLET_COLUMNS = ("N", "P", "R", "T", "V")
SECOND_SLOT_COLUMNS = ("O", "Q", "S", "U")
SWISS_KRONO_TERMS = ("swiss krono", "swisskrono", "свисс кроно", "свис кроно", "krono", "кроно")
DEFAULT_THICKNESS = 16
THICKNESSES = "8|10|12|16|18|19|22|25"
# The form's columns C and D turn these typed numbers into the factory's decor codes.
LETTER_DECORS = {
    "K101": 101,
    "K110": 110,
    "U112": 112,
    "U119": 119,
    "U190": 190,
    "U171": 171,
    "U125": 125,
    "U159": 159,
    "D381": 381,
    "D722": 722,
    "D088": 88,
}
_TYPED_TO_DECOR = {str(number): decor for decor, number in LETTER_DECORS.items()}
# Site names mix in Cyrillic letters that look the same: "D/4075 РЕ".
_LOOKALIKES = str.maketrans("АВЕКМНОРСТХ", "ABEKMHOPCTX")
_NAME_WITH_SPACE = re.compile(r"^\s*D\s*/\s*([0-9A-Za-zА-Яа-я]+)\s+([A-Za-zА-Яа-я]{2})\s*/\s*(.*)$")
_NAME_WITH_SLASHES = re.compile(r"^\s*D\s*/\s*([0-9A-Za-zА-Яа-я]+)\s*/\s*(.*?)\s*/\s*([A-Za-zА-Яа-я]{2})\s*/")
_THICKNESS = re.compile(rf"[*xх×]\s*({THICKNESSES})(?![0-9.,])")
_SIZE_TAIL = re.compile(r"\s*\d+(?:[.,]\d+)?\s*[*xх×]\s*\d.*$")
_CELL = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
_FORMULA = re.compile(r"<f\b[^>]*?/>|<f\b.*?</f>", re.S)
_SHARED_MASTER = re.compile(r'<f\b[^>]*t="shared"[^>]*\bref=')

ORDER_ROW_COLUMNS = [
    "row_id",
    "label",
    "name",
    "decor",
    "structure",
    "thickness",
    "length",
    "pack",
    "index",
    "price_group",
    "needed_sheets",
    "pallets",
    "abc_class",
    "xyz_class",
    "priority",
    "stock_status",
    "available_stock_qty",
    "forecast_qty",
    "product",
    "supplier",
    "low_priority_skipped",
    "in_catalog",
]


@dataclass(frozen=True)
class CatalogEntry:
    pack: int
    index: str
    price_group: object


def template_ready() -> bool:
    return TEMPLATE_PATH.is_file()


def load_catalog() -> tuple[dict[tuple[str, str, int, int], CatalogEntry], float, float]:
    """The factory's material list and its two minimum shares of the price groups."""
    return _read_catalog(str(TEMPLATE_PATH), TEMPLATE_PATH.stat().st_mtime)


@lru_cache(maxsize=1)
def _read_catalog(path: str, _modified: float) -> tuple[dict[tuple[str, str, int, int], CatalogEntry], float, float]:
    return _parse_catalog(path)


def _parse_catalog(source: object) -> tuple[dict[tuple[str, str, int, int], CatalogEntry], float, float]:
    with warnings.catch_warnings():
        # The extensions openpyxl warns about only matter when saving, which never happens here.
        warnings.simplefilter("ignore")
        book = load_workbook(source, data_only=True, read_only=True)
    try:
        if book.worksheets[0].title != ORDER_SHEET:
            raise ValueError(f"Первым должен быть лист «{ORDER_SHEET}».")
        catalog: dict[tuple[str, str, int, int], CatalogEntry] = {}
        for row in book[CATALOG_SHEET].iter_rows(min_row=2, max_col=4, values_only=True):
            parts = str(row[0] or "").split("/")
            # Only boards with one decor on both sides; the form orders the other kind by hand.
            if len(parts) < 9 or parts[1] != parts[2]:
                continue
            thickness = int(float(parts[6].replace(",", ".")))
            length = int(parts[7].upper().split("X")[0])
            key = (parts[1].upper(), parts[3].upper(), thickness, length)
            # Some decors also come in half pallets; the full pallet is the form's default.
            if key in catalog and catalog[key].pack == _default_pack(thickness, length):
                continue
            catalog[key] = CatalogEntry(pack=int(parts[8]), index=str(row[1] or ""), price_group=row[3])
        settings = book[SETTINGS_SHEET]
        return catalog, float(settings["K1"].value or 0), float(settings["K2"].value or 0)
    finally:
        book.close()


def save_template(content: bytes) -> None:
    """Keep the factory's form after checking that orders can be written into it."""
    try:
        catalog = _parse_catalog(BytesIO(content))[0]
        if not catalog:
            raise ValueError("В бланке нет списка материалов.")
        build_order_workbook(pd.DataFrame(columns=ORDER_ROW_COLUMNS), template=content)
    except Exception as error:
        raise ValueError("Это не бланк завода: нужен файл с листом «Заказ на ЛДСП» и строками заказа.") from error
    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = TEMPLATE_PATH.with_name(f"upload-{os.getpid()}.xlsx")
    temporary_path.write_bytes(content)
    temporary_path.replace(TEMPLATE_PATH)


def is_swiss_krono_supplier(value: object) -> bool:
    text = _normalize(value)
    return any(_contains_term(text, term) for term in SWISS_KRONO_TERMS)


def _latin(value: str) -> str:
    return value.upper().translate(_LOOKALIKES)


def _site_decor(raw: str, structure: str, catalog: dict) -> str:
    """The factory's code for a decor: the site writes some of them with a U the factory omits."""
    known = {key[0] for key in catalog if key[1] == structure}
    if raw not in known and re.fullmatch(r"U\d{4,}", raw) and raw[1:] in known:
        return raw[1:]
    return raw


def parse_product(product: str) -> dict[str, object] | None:
    """Decor, structure, thickness and length from a name such as "D/4230 SE/Гикори 2.75*1.83"."""
    match = _NAME_WITH_SPACE.match(product)
    if match:
        raw, structure, name = match.group(1), match.group(2), match.group(3)
    else:
        match = _NAME_WITH_SLASHES.match(product)
        if not match:
            return None
        raw, name, structure = match.group(1), match.group(2), match.group(3)
    catalog = load_catalog()[0]
    structure = _latin(structure)
    thickness = _THICKNESS.search(product.replace(" ", ""))
    return {
        "decor": _site_decor(_latin(raw), structure, catalog),
        "structure": structure,
        "thickness": int(thickness.group(1)) if thickness else DEFAULT_THICKNESS,
        "length": 2800 if "2800" in product or re.search(r"2[.,]8(?!\d)", product) else 2750,
        "name": _SIZE_TAIL.sub("", name).strip() or name.strip(),
    }


def _default_pack(thickness: int, length: int) -> int:
    """Sheets on a pallet by the formula of the form's column X."""
    if thickness == 10:
        return 21 if length == 2750 else 54
    return {12: 40, 16: 33, 18: 29, 19: 25, 22: 24, 25: 21, 8: 67}.get(thickness, 0)


def _pallets_for(sheets: float, pack: int) -> int:
    """To the nearest whole pallet; half a pallet goes up."""
    return math.floor(sheets / pack + 0.5) if pack > 0 and sheets > 0 else 0


def build_order_rows(procurement_forecast: pd.DataFrame, *, only_supplier: bool = True) -> pd.DataFrame:
    """One line per factory material; goods the factory's list does not hold are flagged."""
    if procurement_forecast.empty or "product" not in procurement_forecast.columns:
        return pd.DataFrame(columns=ORDER_ROW_COLUMNS)
    forecast = procurement_forecast
    if only_supplier:
        if "supplier" not in forecast.columns:
            return pd.DataFrame(columns=ORDER_ROW_COLUMNS)
        forecast = forecast[forecast["supplier"].map(is_swiss_krono_supplier)]
    catalog = load_catalog()[0]
    grouped: dict[str, dict[str, object]] = {}
    for record in forecast.to_dict(orient="records"):
        product = str(record["product"])
        parsed = parse_product(product)
        if parsed is None:
            parsed = {"decor": "", "structure": "", "thickness": DEFAULT_THICKNESS, "length": 2750, "name": product}
            row_id = f"?:{product}"
        else:
            row_id = f"{parsed['decor']}/{parsed['structure']}/{parsed['thickness']}/{parsed['length']}"
        entry = catalog.get((parsed["decor"], parsed["structure"], parsed["thickness"], parsed["length"]))
        row = grouped.setdefault(
            row_id,
            {
                "row_id": row_id,
                "label": f"D/{parsed['decor']} {parsed['structure']}" if parsed["decor"] else product,
                "name": parsed["name"],
                "decor": parsed["decor"],
                "structure": parsed["structure"],
                "thickness": parsed["thickness"],
                "length": parsed["length"],
                "pack": entry.pack if entry else _default_pack(parsed["thickness"], parsed["length"]),
                "index": entry.index if entry else "",
                "price_group": entry.price_group if entry else "",
                "needed_sheets": 0.0,
                "pallets": 0,
                "abc_class": "",
                "xyz_class": "",
                "priority": "",
                "stock_status": "",
                "available_stock_qty": 0.0,
                "forecast_qty": 0.0,
                "product": product,
                "supplier": record.get("supplier", ""),
                "low_priority_skipped": False,
                "in_catalog": entry is not None,
            },
        )
        recommended = pd.to_numeric(pd.Series([record.get("recommended_order_qty")]), errors="coerce").fillna(0).iloc[0]
        needed = max(float(recommended), 0.0)
        if _normalize(record.get("priority", "")) == EXCLUDED_ORDER_PRIORITY:
            row["low_priority_skipped"] = bool(row["low_priority_skipped"]) or needed > 0
            needed = 0.0
        row["needed_sheets"] += needed
        for column in ("abc_class", "xyz_class", "priority", "stock_status"):
            row[column] = record.get(column, "") or row[column]
        for column in ("available_stock_qty", "forecast_qty"):
            value = pd.to_numeric(pd.Series([record.get(column)]), errors="coerce").fillna(0).iloc[0]
            row[column] += float(value)
    rows = pd.DataFrame(list(grouped.values()), columns=ORDER_ROW_COLUMNS)
    rows["pallets"] = [
        _pallets_for(sheets, pack) if known else 0
        for sheets, pack, known in zip(rows["needed_sheets"], rows["pack"], rows["in_catalog"])
    ]
    return rows.sort_values(["pallets", "needed_sheets"], ascending=False, kind="stable").reset_index(drop=True)


def rows_from_saved_order(items: pd.DataFrame) -> pd.DataFrame:
    """The form for a saved order: its own quantities, not the current recommendation."""
    return build_order_rows(items.rename(columns={"quantity": "recommended_order_qty"}), only_supplier=False)


def _ordered(order_rows: pd.DataFrame) -> pd.DataFrame:
    return order_rows[order_rows["in_catalog"].astype(bool) & pd.to_numeric(order_rows["pallets"]).gt(0)]


def order_items_from_rows(order_rows: pd.DataFrame) -> pd.DataFrame:
    """Order lines in sheets, the unit the forecast counts in, for goods the factory can supply."""
    ordered = _ordered(order_rows)
    return pd.DataFrame(
        {
            "product": ordered["product"],
            "quantity": (ordered["pallets"] * ordered["pack"]).astype(float),
            "unit_cost": 0.0,
            "comment": "",
        },
        columns=["product", "quantity", "unit_cost", "comment"],
    ).reset_index(drop=True)


def order_totals(order_rows: pd.DataFrame) -> dict[str, float]:
    """Totals as the form computes them, with the two price group shares of its top row."""
    ordered = _ordered(order_rows)
    sheets = ordered["pallets"] * ordered["pack"]
    total = float(sheets.sum())
    group_sheets = {group: float(sheets[ordered["price_group"].eq(group)].sum()) for group in range(1, 8)}
    return {
        "pallets": float(ordered["pallets"].sum()),
        "sheets": total,
        "area": total * SHEET_AREA_M2,
        "middle_share": (group_sheets[3] + group_sheets[4]) / total if total else 0.0,
        "top_share": (group_sheets[5] + group_sheets[6] + group_sheets[7]) / total if total else 0.0,
        **{f"group_{group}": value for group, value in group_sheets.items()},
    }


class _Plain:
    """A value that replaces the cell's formula: half pallets are typed over column X, as the form asks."""

    def __init__(self, value: object) -> None:
        self.value = value


class _Formula:
    """A cell that must hold this formula, whatever the blank has there."""

    def __init__(self, formula: str, value: object) -> None:
        self.formula = formula
        self.value = value


class _Error:
    def __init__(self, code: str) -> None:
        self.code = code


def _render_cell(reference: str, attributes: str, inner: str | None, new: object) -> str:
    formula_match = _FORMULA.search(inner or "")
    formula = formula_match.group(0) if formula_match else ""
    if isinstance(new, _Plain):
        new, formula = new.value, ""
    elif isinstance(new, _Formula):
        formula, new = f"<f>{escape(new.formula)}</f>", new.value
    attributes = re.sub(r'\s+t="[^"]*"', "", attributes)
    if isinstance(new, _Error):
        return f'<c r="{reference}"{attributes} t="e">{formula}<v>{new.code}</v></c>'
    if new is None or new == "":
        if formula:
            return f'<c r="{reference}"{attributes} t="str">{formula}<v></v></c>' if new == "" else f'<c r="{reference}"{attributes}>{formula}</c>'
        return f'<c r="{reference}"{attributes}/>'
    if isinstance(new, str):
        if formula:
            return f'<c r="{reference}"{attributes} t="str">{formula}<v>{escape(new)}</v></c>'
        return f'<c r="{reference}"{attributes} t="inlineStr"><is><t>{escape(new)}</t></is></c>'
    return f'<c r="{reference}"{attributes}>{formula}<v>{new!r}</v></c>'


def _typed_decor(decor: str) -> object:
    if decor in LETTER_DECORS:
        return LETTER_DECORS[decor]
    return int(decor) if decor.isdigit() else decor


def _shown_decor(decor: str) -> object:
    return int(decor) if decor.isdigit() else decor


def build_order_workbook(order_rows: pd.DataFrame, *, template: bytes | None = None) -> bytes:
    """The factory's form with the ordered goods in the first half of the month."""
    ordered = _ordered(order_rows)
    if template is None:
        template = TEMPLATE_PATH.read_bytes()
    with ZipFile(BytesIO(template)) as archive:
        sheet_xml = archive.read(ORDER_SHEET_PART).decode("utf-8")
    cells = {match.group(1) + match.group(2): match for match in _CELL.finditer(sheet_xml)}

    updates: dict[str, object] = {}
    for row in range(FIRST_ROW, LAST_ROW + 1):
        for column in ("E", "F", *PALLET_COLUMNS):
            updates[f"{column}{row}"] = None
        for column in ("M", *SECOND_SLOT_COLUMNS, "W", "Y"):
            updates[f"{column}{row}"] = 0
        updates[f"C{row}"] = 0
        updates[f"D{row}"] = ""
        updates[f"L{row}"] = _Error("#N/A")
        updates[f"AC{row}"] = _Error("#N/A")
        # The blank came with a number typed over this formula, which inflated its total.
        if "<f" not in (cells[f"W{row}"].group(4) or ""):
            updates[f"W{row}"] = _Formula(f"M{row}+O{row}+Q{row}+S{row}+U{row}", 0)

    row = FIRST_ROW
    for record in ordered.to_dict(orient="records"):
        pack = int(record["pack"])
        needs_typed_pack = pack != _default_pack(int(record["thickness"]), int(record["length"]))
        # A half pallet is typed over column X, which must not hold the master of a shared formula.
        while needs_typed_pack and row <= LAST_ROW and _SHARED_MASTER.search(cells[f"X{row}"].group(4) or ""):
            row += 1
        if row > LAST_ROW:
            raise ValueError(f"В бланке завода {LAST_ROW - FIRST_ROW + 1} строк, а позиций в заказе больше.")
        decor, structure = str(record["decor"]), str(record["structure"])
        thickness, length, pallets = int(record["thickness"]), int(record["length"]), int(record["pallets"])
        sheets = pack * pallets
        size = "2750x1830" if length == 2750 else "2800x2070"
        shown_thickness = f"0{thickness}" if thickness < 10 else str(thickness)
        price_group = record["price_group"]
        updates.update(
            {
                f"E{row}": _typed_decor(decor),
                f"F{row}": _typed_decor(decor),
                f"C{row}": _shown_decor(decor),
                f"D{row}": _shown_decor(decor),
                f"G{row}": structure,
                f"H{row}": structure,
                f"I{row}": thickness,
                f"J{row}": length,
                f"K{row}": size,
                f"L{row}": str(record["index"]),
                f"M{row}": sheets,
                f"N{row}": pallets,
                f"W{row}": updates[f"W{row}"] if isinstance(updates[f"W{row}"], _Formula) else sheets,
                f"X{row}": _Plain(pack) if needs_typed_pack else pack,
                f"Y{row}": sheets * SHEET_AREA_M2,
                f"Z{row}": f"0{pack}",
                f"AA{row}": f"D/{decor}/{decor}/{structure}/{structure}//{shown_thickness},0/{size}/0{pack}/1",
                f"AC{row}": price_group if price_group not in (None, "") else _Error("#N/A"),
            }
        )
        if isinstance(updates[f"W{row}"], _Formula):
            updates[f"W{row}"].value = sheets
        row += 1

    totals = order_totals(order_rows)
    updates.update(
        {
            f"M{TOTAL_ROW}": totals["sheets"],
            f"N{TOTAL_ROW}": totals["pallets"],
            f"W{TOTAL_ROW}": totals["sheets"],
            f"Y{TOTAL_ROW}": totals["area"],
            "N1": totals["middle_share"],
            "P1": totals["top_share"],
        }
    )
    for column in (*SECOND_SLOT_COLUMNS, *PALLET_COLUMNS[1:]):
        updates[f"{column}{TOTAL_ROW}"] = 0
    for group, column in zip(range(1, 8), ("AD", "AE", "AF", "AG", "AH", "AI", "AJ")):
        updates[f"{column}2"] = totals[f"group_{group}"]

    def replace(match: re.Match[str]) -> str:
        reference = match.group(1) + match.group(2)
        if reference not in updates:
            return match.group(0)
        return _render_cell(reference, match.group(3), match.group(4), updates[reference])

    new_sheet = _CELL.sub(replace, sheet_xml)
    target = BytesIO()
    with ZipFile(BytesIO(template)) as archive, ZipFile(target, "w", ZIP_DEFLATED) as output:
        for item in archive.infolist():
            # Excel rebuilds the calculation chain; a stale one makes it ask to repair the file.
            if item.filename == "xl/calcChain.xml":
                continue
            data = archive.read(item.filename)
            if item.filename == ORDER_SHEET_PART:
                data = new_sheet.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                data = re.sub(rb'<Override [^>]*calcChain[^>]*/>', b"", data)
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = re.sub(rb'<Relationship [^>]*calcChain[^>]*/>', b"", data)
            elif item.filename == "xl/workbook.xml":
                data = re.sub(rb"<calcPr\b([^>]*?)/>", rb'<calcPr\1 fullCalcOnLoad="1"/>', data)
            output.writestr(item, data)
    return target.getvalue()


@dataclass
class ParsedFactoryFile:
    """Final pallets from the form the user edited and uploaded back."""

    fingerprint: str = ""
    pallets: dict[str, float] = field(default_factory=dict)
    unknown: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _file_number(value: object) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0.0
    try:
        number = float(str(value).replace(",", ".").replace(" ", "")) if isinstance(value, str) else float(value)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(number) or number < 0 else number


def parse_order_file(content: bytes, order_rows: pd.DataFrame) -> ParsedFactoryFile:
    parsed = ParsedFactoryFile(fingerprint=hashlib.sha256(content).hexdigest()[:12])
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            book = load_workbook(BytesIO(content), data_only=True)
        sheet = book[ORDER_SHEET]
    except Exception:
        parsed.errors.append("Не удалось прочитать файл. Нужен бланк завода Swiss Krono с листом «Заказ на ЛДСП».")
        return parsed
    known = set(order_rows["row_id"])
    for row in range(FIRST_ROW, LAST_ROW + 1):
        typed = sheet[f"E{row}"].value
        slots = [_file_number(sheet[f"{column}{row}"].value) for column in PALLET_COLUMNS]
        if typed in (None, ""):
            continue
        if None in slots:
            parsed.errors.append(f"Строка {row}, декор {typed}: количество паллет должно быть числом не меньше нуля.")
            continue
        pallets = sum(slots)
        if pallets <= 0:
            continue
        typed_text = str(int(typed)) if isinstance(typed, float) and typed.is_integer() else str(typed).strip()
        decor = _TYPED_TO_DECOR.get(typed_text, _latin(typed_text))
        structure = _latin(str(sheet[f"G{row}"].value or "").strip())
        thickness = _file_number(sheet[f"I{row}"].value)
        length = _file_number(sheet[f"J{row}"].value)
        row_id = f"{decor}/{structure}/{int(thickness or 0)}/{int(length or 0)}"
        if row_id in known:
            parsed.pallets[row_id] = parsed.pallets.get(row_id, 0.0) + pallets
        else:
            parsed.unknown.append(f"D/{decor} {structure} {int(thickness or 0)} мм")
    return parsed


def apply_order_file(order_rows: pd.DataFrame, parsed: ParsedFactoryFile) -> pd.DataFrame:
    """The lines with the file's pallets: a line left blank in the file is not ordered."""
    result = order_rows.copy()
    result["pallets"] = [parsed.pallets.get(row_id, 0.0) for row_id in result["row_id"]]
    return result


def order_file_comment(filename: str, parsed: ParsedFactoryFile) -> str:
    return f"Бланк завода Swiss Krono · файл {filename} · {parsed.fingerprint}"
