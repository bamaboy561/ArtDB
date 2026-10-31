"""Order forms in the layout a supplier expects, filled from the procurement forecast."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import math
import re
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


SECTION_BOARDS = "boards"
SECTION_HPL = "hpl"


@dataclass(frozen=True)
class OrderFormRow:
    code: str
    name: str
    old_code: str = ""
    section: str = SECTION_BOARDS
    # A generic name ("белый") must not pull in unrelated goods.
    match_by_name: bool = True


WQ_BOARD_FORMAT = "2800*1220*18мм"
WQ_ORDER_ROWS: tuple[OrderFormRow, ...] = (
    OrderFormRow("WQ63815M", "Ванильный крем", "US77102/PGR4102"),
    OrderFormRow("WQ63814M", "Молочный пломбир", "US77103/PGR4103"),
    OrderFormRow("WQ63893M", "Маффин", "US77301/PGR4201"),
    OrderFormRow("WQ63845M", "Брауни", "US77302/PGR4202"),
    OrderFormRow("US77304/PGR4204", "Орео", "US77304/PGR4204"),
    OrderFormRow("US77305/PGR4205", "Тирамису", "US77305/PGR4205"),
    OrderFormRow("WQ63864M", "Песочное печенье", "US77202/PGR4302"),
    OrderFormRow("WQ63819M", "Бисквит", "US77306/PGR4206"),
    OrderFormRow("WQ63809M", "Вупи пай", "US77307/PGR4207"),
    OrderFormRow("WQ63812M", "Чизкейк", "US77311/PGR4211"),
    OrderFormRow("WQ63853M", "Меренга", "US77312/PGR4212"),
    OrderFormRow("WQ63811M", "Глазурь", "US77314/PGR4214"),
    OrderFormRow("WQ63894M", "Пряный апельсин", "US77401/PGR4401"),
    OrderFormRow("WQ63851M", "Спелый авокадо", "US77601/PGR4601"),
    OrderFormRow("WQ63167M", "Мятный кекс", "US77602/PGR4602"),
    OrderFormRow("WQ63102M-1", "Ментоловый лёд", "US77603/PGR4603"),
    OrderFormRow("WQ63895M", "Бабл гам", "US77702/PGR4702"),
    OrderFormRow("WQ63850M", "Черничный тарт"),
    OrderFormRow("WQ63098M-1", "Панна кота"),
    OrderFormRow("WQ63839M", "Пьяная вишня"),
    OrderFormRow("WQ63833M", "Моти", "NEW"),
    OrderFormRow("WQ63033M-1", "Туман", "NEW"),
    OrderFormRow("WQ63142LJ", "Титан", "NEW"),
    OrderFormRow("WQ63835M", "Крем Брюле", "NEW"),
    OrderFormRow("WQ63898MJ", "Бронза", "NEW"),
    OrderFormRow("WQ63825MJ", "Бархат", "NEW"),
    OrderFormRow("WQWE71406-2(1386)", "Кассела кор", "NEW"),
    OrderFormRow("WQWE71406-1(1385)", "Кассела натур", "NEW"),
    OrderFormRow("WQWE7140-6-(1367)", "Кассела светл", "NEW"),
    OrderFormRow("WQS31008-MT", "Изысканный Камень", section=SECTION_HPL),
    OrderFormRow("WQS31015-MT", "Мрамор Венато", section=SECTION_HPL),
    OrderFormRow("81050", "белый", section=SECTION_HPL, match_by_name=False),
    OrderFormRow("WQS1837-3 FR", "Молочный Кварцит", section=SECTION_HPL),
    OrderFormRow("WQW47183-1 ELM", "Орех Восточный", section=SECTION_HPL),
    OrderFormRow("WQW47183-2 ELM", "Дерево Песчаное", section=SECTION_HPL),
)

ORDER_FORM_COLUMNS = [
    "code",
    "quantity",
    "old_code",
    "edge_or_sheets",
    "name",
    "section",
    "abc_class",
    "xyz_class",
    "priority",
    "stock_status",
    "available_stock_qty",
    "forecast_qty",
    "found",
    "product",
    "edge_product",
    "supplier",
]

_EDGE_SIZE = re.compile(r"(?<![0-9])1\s*[*xх×]\s*22(?![0-9])")
_BOARD_SIZE_PARTS = ("18", "1220", "2800")
_NUMBER = re.compile(r"[0-9]+")


def _normalize(value: object) -> str:
    return " ".join(str(value or "").casefold().replace("ё", "е").split())


def _contains_term(text: str, term: str) -> bool:
    term = _normalize(term)
    return bool(term) and re.search(rf"(?<![0-9a-zа-я]){re.escape(term)}(?![0-9a-zа-я])", text) is not None


def _is_edge(text: str) -> bool:
    return _EDGE_SIZE.search(text) is not None


def _is_board(text: str) -> bool:
    numbers = set(_NUMBER.findall(text))
    return all(part in numbers for part in _BOARD_SIZE_PARTS)


def _old_code_terms(old_code: str) -> list[str]:
    return [part for part in old_code.split("/") if part and part != "NEW"]


def _row_mask(row: OrderFormRow, texts: pd.Series, kinds: pd.Series, kind: str) -> pd.Series:
    """Rows of the forecast that are this form row: by article first, by name otherwise."""
    candidates = texts[kinds.eq(kind)]
    terms = [row.code, *_old_code_terms(row.old_code)]
    by_code = candidates[candidates.map(lambda text: any(_contains_term(text, term) for term in terms)).astype(bool)]
    if by_code.empty and row.match_by_name:
        by_code = candidates[candidates.map(lambda text: _contains_term(text, row.name)).astype(bool)]
    return texts.index.isin(by_code.index)


def _product_kind(text: str) -> str:
    if _is_edge(text):
        return "edge"
    if "hpl" in text:
        return "hpl"
    if _is_board(text):
        return "board"
    return ""


def _whole_quantity(values: pd.Series) -> int:
    total = float(pd.to_numeric(values, errors="coerce").fillna(0).clip(lower=0).sum())
    return math.ceil(total - 1e-9) if total > 0 else 0


def build_order_form_rows(
    procurement_forecast: pd.DataFrame,
    rows: tuple[OrderFormRow, ...] = WQ_ORDER_ROWS,
) -> pd.DataFrame:
    """One line per form row: boards, edge metres or HPL sheets, as the forecast recommends."""
    if procurement_forecast.empty or "product" not in procurement_forecast.columns:
        forecast = pd.DataFrame(columns=["product", "recommended_order_qty"])
    else:
        forecast = procurement_forecast.reset_index(drop=True)
    texts = forecast["product"].map(_normalize)
    kinds = texts.map(_product_kind)

    def first_value(frame: pd.DataFrame, column: str, default: object) -> object:
        if column not in frame.columns or frame.empty:
            return default
        return frame[column].iloc[0]

    def total(frame: pd.DataFrame, column: str) -> float:
        if column not in frame.columns:
            return 0.0
        return float(pd.to_numeric(frame[column], errors="coerce").fillna(0).sum())

    def first_product(frame: pd.DataFrame) -> str:
        return str(frame["product"].iloc[0]) if not frame.empty else ""

    records = []
    for row in rows:
        main_kind = "hpl" if row.section == SECTION_HPL else "board"
        main = forecast[_row_mask(row, texts, kinds, main_kind)]
        main_quantity = _whole_quantity(main["recommended_order_qty"]) if not main.empty else 0
        edge_found = False
        edge = forecast.iloc[0:0]
        if row.section == SECTION_HPL:
            quantity, edge_or_sheets = 0, main_quantity
        else:
            edge = forecast[_row_mask(row, texts, kinds, "edge")]
            edge_found = not edge.empty
            quantity = main_quantity
            edge_or_sheets = _whole_quantity(edge["recommended_order_qty"]) if edge_found else 0
        records.append(
            {
                "code": row.code,
                "quantity": quantity,
                "old_code": row.old_code,
                "edge_or_sheets": edge_or_sheets,
                "name": row.name,
                "section": row.section,
                "abc_class": first_value(main, "abc_class", ""),
                "xyz_class": first_value(main, "xyz_class", ""),
                "priority": first_value(main, "priority", ""),
                "stock_status": first_value(main, "stock_status", ""),
                "available_stock_qty": total(main, "available_stock_qty"),
                "forecast_qty": total(main, "forecast_qty"),
                "found": (not main.empty) or edge_found,
                "product": first_product(main),
                "edge_product": first_product(edge),
                "supplier": first_value(main, "supplier", ""),
            }
        )
    return pd.DataFrame(records, columns=ORDER_FORM_COLUMNS)


def order_form_from_items(items: pd.DataFrame) -> pd.DataFrame:
    """The form for a saved order: its own quantities, not the current recommendation."""
    return build_order_form_rows(items.rename(columns={"quantity": "recommended_order_qty"}))


def order_items_from_form(order_rows: pd.DataFrame) -> pd.DataFrame:
    """Order lines for the form as it stands, so the ordered goods count as in transit."""
    quantities: dict[str, float] = {}
    for record in order_rows.to_dict(orient="records"):
        if record["section"] == SECTION_HPL:
            lines = ((record["product"], record["edge_or_sheets"]),)
        else:
            lines = ((record["product"], record["quantity"]), (record["edge_product"], record["edge_or_sheets"]))
        for product, quantity in lines:
            if product and float(quantity) > 0:
                quantities[product] = quantities.get(product, 0.0) + float(quantity)
    return pd.DataFrame(
        [
            {"product": product, "quantity": quantity, "unit_cost": 0.0, "comment": ""}
            for product, quantity in quantities.items()
        ],
        columns=["product", "quantity", "unit_cost", "comment"],
    )


def order_form_supplier(order_rows: pd.DataFrame, default: str = "WQ") -> str:
    """The site's name for this supplier: the one most of the form's goods are assigned to."""
    suppliers = order_rows["supplier"].fillna("").astype(str).str.strip()
    suppliers = suppliers[suppliers.ne("") & suppliers.str.casefold().ne("не назначен")]
    return str(suppliers.mode().iloc[0]) if not suppliers.empty else default


def _store_total_value(content: bytes, total_cell: str, total_value: int) -> bytes:
    """Save the computed total next to its formula: Excel's Protected View shows only stored values."""
    source = BytesIO(content)
    target = BytesIO()
    pattern = re.compile(rf'(<c r="{total_cell}"[^>]*><f>[^<]*</f>)<v\s*/>|(<c r="{total_cell}"[^>]*><f>[^<]*</f>)<v></v>')
    with ZipFile(source) as archive, ZipFile(target, "w", ZIP_DEFLATED) as output:
        for item in archive.infolist():
            data = archive.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                text = data.decode("utf-8")
                text = pattern.sub(lambda m: f"{m.group(1) or m.group(2)}<v>{total_value}</v>", text)
                data = text.encode("utf-8")
            output.writestr(item, data)
    return target.getvalue()


def build_wq_order_workbook(order_rows: pd.DataFrame) -> bytes:
    """The WQ form: boards with edge metres, then the HPL block with sheets, then the boards total."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"

    medium = Side(style="medium")
    thin = Side(style="thin")
    yellow = PatternFill("solid", fgColor="FFFF00")
    blue = PatternFill("solid", fgColor="B4C6E7")
    centered = Alignment(horizontal="center", vertical="center")

    for column, width in {"A": 3.7, "B": 21.0, "C": 17.9, "D": 18.3, "E": 17.6, "F": 20.7}.items():
        sheet.column_dimensions[column].width = width

    sheet["C3"] = WQ_BOARD_FORMAT
    sheet["C3"].font = Font(name="Calibri", size=11, bold=True)
    sheet["C3"].alignment = centered

    for column, title in enumerate(("Код", "кол-во", "код старый", "кромка", "Название"), start=2):
        cell = sheet.cell(row=4, column=column, value=title)
        cell.fill = yellow
        cell.font = Font(name="Calibri", size=12, bold=True)
        cell.alignment = centered
        cell.border = Border(left=medium, right=medium, top=medium, bottom=medium)

    def write_row(row_number: int, record: dict[str, object], *, shaded_code: bool) -> None:
        values = (
            record["code"],
            int(record["quantity"]) or None,
            record["old_code"] or None,
            int(record["edge_or_sheets"]) or None,
            record["name"],
        )
        for column, value in enumerate(values, start=2):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.font = Font(name="Calibri", size=11, bold=column == 2)
            cell.alignment = Alignment(horizontal=None if column == 2 else "center", vertical="center")
            cell.border = Border(
                left=medium if column == 2 else thin,
                right=medium if column == 6 else thin,
                bottom=thin,
            )
        if shaded_code:
            sheet.cell(row=row_number, column=2).fill = blue

    records = order_rows.to_dict(orient="records")
    row_number = 5
    for record in records:
        if record["section"] == SECTION_BOARDS:
            write_row(row_number, record, shaded_code=True)
            row_number += 1

    hpl_records = [record for record in records if record["section"] == SECTION_HPL]
    if hpl_records:
        for column in range(2, 7):
            cell = sheet.cell(row=row_number, column=column)
            cell.fill = blue
            cell.border = Border(
                left=medium if column == 2 else None,
                right=medium if column == 6 else None,
                top=medium,
                bottom=thin,
            )
        title = sheet.cell(row=row_number, column=2, value="HPL")
        title.font = Font(name="Calibri", size=16, bold=True)
        title.alignment = centered
        sheet.row_dimensions[row_number].height = 21
        sheet.cell(row=row_number, column=5, value="листов").alignment = centered
        sheet.cell(row=row_number, column=5).font = Font(name="Calibri", size=11, bold=True)
        row_number += 1
        for record in hpl_records:
            write_row(row_number, record, shaded_code=False)
            row_number += 1

    total_cell = f"C{row_number}"
    sheet[total_cell] = f"=SUM(C5:C{row_number - 1})"
    sheet[total_cell].font = Font(name="Calibri", size=14, bold=True)
    sheet[total_cell].alignment = centered
    sheet.row_dimensions[row_number].height = 19.5

    sheet.print_area = f"B3:F{row_number}"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = BytesIO()
    workbook.save(buffer)
    return _store_total_value(buffer.getvalue(), total_cell, int(order_rows["quantity"].sum()))
